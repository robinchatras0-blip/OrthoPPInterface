"""Unit tests validating publication-grade figures, styled Excel, and interactive HTML report generation.

Verifies:
1. Generation of SUMMARY.html, SUMMARY.xlsx and all 6 individual figures in figures/ (both .svg and .png at 300 DPI).
2. Proper inclusion of inline SVGs, executive takeaways, podium cards, and filtering JS in HTML.
3. Robustness on edge cases (empty DataFrame, single-row, missing optional energy/scaffold columns).
"""

import os
import openpyxl
import pandas as pd
import pytest

from export_styled_reports import (
    generate_interactive_html,
    generate_styled_excel,
    plot_cross_reactivity_matrix,
    plot_specificity_landscape,
    plot_interaction_triad_distributions,
    plot_structural_funnel,
    plot_induced_fit_and_biophysics,
    plot_pipeline_attrition_funnel,
)


@pytest.fixture
def sample_full_cohort_csv(tmp_path):
    """Create a realistic cohort CSV with full RF3, ipSAE, PyRosetta, and scaffold metrics."""
    csv_file = tmp_path / "orthogonality_scores.csv"
    data = [
        {
            "design_id": "Des_01_Champion", "parent_a_motif": "motif_alpha", "folding_engine": "rf3",
            "status": "ok", "iptm_rescue": 0.88, "iptm_negative": 0.18, "iptm_rupture": 0.22,
            "iptm_ceiling": 0.84, "iptm_rescue_rel": 1.05, "f_ortho_iptm": 0.35, "f_iptm_rel": 0.38,
            "ipsae_rescue": 0.85, "ipsae_rupture": 0.20, "ipsae_negative": 0.19, "ipsae_ceiling": 0.82, "f_ipsae": 0.36,
            "pass_rescue": True, "pass_rupture": True, "pass_negative": True, "pass_dockq": True,
            "dockq": 0.78, "dockq_quality": "High", "fnat": 0.82, "irms": 0.95, "lrms": 1.45,
            "selfcons_lrms": 1.20, "plddt_rescue": 89.5, "n_mut_A": 3, "n_mut_B": 4,
            "scaffold_flex_rmsd": 1.85, "scaffold_is_native": False,
            "dG_rescue": -35.4, "dG_negative": -5.2, "dG_rupture": -4.1, "dG_native": -32.0,
            "dsasa_rescue": 1650.0, "unsat_hb_rescue": 2.0, "f_energy": 0.42, "pass_energy": True,
            "f_ortho": 0.39, "passes": True,
        },
        {
            "design_id": "Des_02_Orthogonal", "parent_a_motif": "motif_beta", "folding_engine": "rf3",
            "status": "ok", "iptm_rescue": 0.74, "iptm_negative": 0.28, "iptm_rupture": 0.29,
            "iptm_ceiling": 0.84, "iptm_rescue_rel": 0.88, "f_ortho_iptm": 0.18, "f_iptm_rel": 0.19,
            "ipsae_rescue": 0.72, "ipsae_rupture": 0.27, "ipsae_negative": 0.26, "ipsae_ceiling": 0.82, "f_ipsae": 0.18,
            "pass_rescue": True, "pass_rupture": True, "pass_negative": True, "pass_dockq": True,
            "dockq": 0.55, "dockq_quality": "Medium", "fnat": 0.65, "irms": 1.45, "lrms": 3.10,
            "selfcons_lrms": 2.80, "plddt_rescue": 83.2, "n_mut_A": 2, "n_mut_B": 3,
            "scaffold_flex_rmsd": 0.45, "scaffold_is_native": True,
            "dG_rescue": -28.1, "dG_negative": -8.5, "dG_rupture": -9.0, "dG_native": -32.0,
            "dsasa_rescue": 1420.0, "unsat_hb_rescue": 4.0, "f_energy": 0.22, "pass_energy": True,
            "f_ortho": 0.20, "passes": True,
        },
        {
            "design_id": "Des_03_CrossAffinity", "parent_a_motif": "motif_alpha", "folding_engine": "rf3",
            "status": "ok", "iptm_rescue": 0.68, "iptm_negative": 0.52, "iptm_rupture": 0.48,
            "iptm_ceiling": 0.84, "iptm_rescue_rel": 0.81, "f_ortho_iptm": -0.05, "f_iptm_rel": -0.04,
            "ipsae_rescue": 0.65, "ipsae_rupture": 0.49, "ipsae_negative": 0.50, "ipsae_ceiling": 0.82, "f_ipsae": -0.05,
            "pass_rescue": True, "pass_rupture": False, "pass_negative": False, "pass_dockq": True,
            "dockq": 0.52, "dockq_quality": "Medium", "fnat": 0.58, "irms": 1.60, "lrms": 3.40,
            "selfcons_lrms": 3.10, "plddt_rescue": 80.0, "n_mut_A": 1, "n_mut_B": 2,
            "scaffold_flex_rmsd": 2.10, "scaffold_is_native": False,
            "dG_rescue": -22.0, "dG_negative": -20.5, "dG_rupture": -19.0, "dG_native": -32.0,
            "dsasa_rescue": 1300.0, "unsat_hb_rescue": 6.0, "f_energy": -0.02, "pass_energy": False,
            "f_ortho": -0.03, "passes": False,
        },
        {
            "design_id": "Des_04_OffTarget", "parent_a_motif": "motif_gamma", "folding_engine": "rf3",
            "status": "early_exit_off_target", "iptm_rescue": 0.45, "iptm_negative": 0.30, "iptm_rupture": 0.25,
            "iptm_ceiling": 0.84, "iptm_rescue_rel": 0.53, "f_ortho_iptm": 0.05, "f_iptm_rel": 0.04,
            "ipsae_rescue": 0.42, "ipsae_rupture": 0.25, "ipsae_negative": 0.28, "ipsae_ceiling": 0.82, "f_ipsae": 0.03,
            "pass_rescue": False, "pass_rupture": True, "pass_negative": True, "pass_dockq": False,
            "dockq": 0.12, "dockq_quality": "Incorrect", "fnat": 0.05, "irms": 4.50, "lrms": 15.20,
            "selfcons_lrms": 14.80, "plddt_rescue": 72.0, "n_mut_A": 4, "n_mut_B": 4,
            "scaffold_flex_rmsd": 3.80, "scaffold_is_native": False,
            "dG_rescue": -10.5, "dG_negative": -8.0, "dG_rupture": -7.5, "dG_native": -32.0,
            "dsasa_rescue": 850.0, "unsat_hb_rescue": 9.0, "f_energy": 0.01, "pass_energy": False,
            "f_ortho": 0.02, "passes": False,
        },
    ]
    df = pd.DataFrame(data)
    df.to_csv(csv_file, index=False)
    return csv_file


def test_complete_report_and_figures_generation(sample_full_cohort_csv, tmp_path):
    """Verify SUMMARY.html, SUMMARY.xlsx, and individual figure files in figures/ are created."""
    html_file = tmp_path / "SUMMARY.html"
    xlsx_file = tmp_path / "SUMMARY.xlsx"

    generate_styled_excel(str(sample_full_cohort_csv), str(xlsx_file))
    generate_interactive_html(str(sample_full_cohort_csv), str(html_file))

    # 1. Output files exist
    assert html_file.exists(), "SUMMARY.html must be generated"
    assert xlsx_file.exists(), "SUMMARY.xlsx must be generated"
    assert html_file.stat().st_size > 5000, "SUMMARY.html must contain substantial dashboard content"
    assert xlsx_file.stat().st_size > 1000, "SUMMARY.xlsx must not be empty"

    # 2. figures/ directory exists and contains all 6 figures in both .svg and .png formats
    fig_dir = tmp_path / "figures"
    assert fig_dir.exists() and fig_dir.is_dir(), "Dossier figures/ must be created"

    expected_figures = [
        "figure1_cross_reactivity",
        "figure2_specificity_landscape",
        "figure3_triad_distributions",
        "figure4_structural_funnel",
        "figure5_induced_fit_biophysics",
        "figure6_attrition_funnel",
    ]

    for fig_base in expected_figures:
        svg_path = fig_dir / f"{fig_base}.svg"
        png_path = fig_dir / f"{fig_base}.png"
        assert svg_path.exists(), f"Figure {fig_base}.svg must exist in figures/"
        assert png_path.exists(), f"Figure {fig_base}.png must exist in figures/"
        assert svg_path.stat().st_size > 500, f"Figure {fig_base}.svg must not be empty"
        assert png_path.stat().st_size > 1000, f"Figure {fig_base}.png must not be empty"

    # 3. HTML Content checks
    html_content = html_file.read_text(encoding="utf-8")

    # All 6 SVGs are embedded inline in HTML
    assert html_content.count("<svg") >= 6, "HTML must contain at least 6 inline SVG figures"

    # Executive key takeaways and metrics
    assert "Diagnostic Summary &amp; Key Takeaways" in html_content or "Diagnostic Summary" in html_content
    assert "2 / 4 designs (50.0%)" in html_content
    assert "Des_01_Champion" in html_content

    # Podium cards
    assert "Top 3 Champions Podium" in html_content
    assert "🥇 1st - Champion" in html_content

    # Quick filters & Interactive controls
    assert "setQuickFilter('all'" in html_content
    assert "setQuickFilter('passed'" in html_content
    assert "sortTable(" in html_content

    # 4. Excel Workbook checks
    wb = openpyxl.load_workbook(str(xlsx_file))
    assert "Executive Summary" in wb.sheetnames
    assert "Raw All-Atom Metrics" in wb.sheetnames

    ws1 = wb["Executive Summary"]
    assert ws1["A1"].value == "ORTHOPPINTERFACE - PROTEIN-PROTEIN ORTHOGONALITY REPORT"
    assert ws1["K5"].value == "CHAMPION ORTHOGONAL"
    assert ws1["K6"].value == "ORTHOGONAL"
    assert ws1["K7"].value == "NON ORTHOGONAL"
    assert "OFF-TARGET" in str(ws1["K8"].value)


def test_edge_case_empty_dataframe(tmp_path):
    """Verify that export functions handle an empty DataFrame without raising exceptions."""
    empty_csv = tmp_path / "empty.csv"
    pd.DataFrame().to_csv(empty_csv, index=False)

    html_file = tmp_path / "empty.html"
    xlsx_file = tmp_path / "empty.xlsx"

    # Should not crash
    generate_styled_excel(str(empty_csv), str(xlsx_file))
    generate_interactive_html(str(empty_csv), str(html_file))

    assert html_file.exists()
    assert xlsx_file.exists()


def test_edge_case_single_minimal_row(tmp_path):
    """Verify single-row dataframe lacking optional PyRosetta and ipSAE columns."""
    minimal_csv = tmp_path / "minimal.csv"
    df = pd.DataFrame([{
        "design_id": "solo_candidate",
        "parent_a_motif": "solo_motif",
        "iptm_rescue": 0.75,
        "iptm_negative": 0.25,
        "iptm_rupture": 0.20,
        "f_ortho": 0.30,
        "dockq": 0.65,
        "passes": True,
    }])
    df.to_csv(minimal_csv, index=False)

    html_file = tmp_path / "minimal.html"
    xlsx_file = tmp_path / "minimal.xlsx"

    generate_styled_excel(str(minimal_csv), str(xlsx_file))
    generate_interactive_html(str(minimal_csv), str(html_file))

    assert html_file.exists()
    assert xlsx_file.exists()
    html_text = html_file.read_text(encoding="utf-8")
    assert "solo_candidate" in html_text
    assert "<svg" in html_text
