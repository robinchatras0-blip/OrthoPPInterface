import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_styled_excel(csv_path, xlsx_path):
    if not os.path.exists(csv_path):
        return

    df = pd.read_csv(csv_path)

    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Sheet 1: Executive Orthogonality Dashboard
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:K1")
    title_cell = ws1["A1"]
    title_cell.value = "ORTHOPPINTERFACE - RAPPORT D'ORTHOGONALITÉ PROTÉINE-PROTÉINE"
    title_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 35

    # Subtitle Banner
    ws1.merge_cells("A2:K2")
    sub_cell = ws1["A2"]
    sub_cell.value = "Évaluation Haute Résolution RoseTTAFold-3 All-Atom (RF3) & Conformité DockQ CAPRI"
    sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
    sub_cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # Columns for Executive Summary
    cols_summary = [
        ("design_id", "Paire Orthogonale (A'·B')", 40),
        ("parent_a_motif", "Motif A' Parent", 24),
        ("iptm_rescue", "iPTM Sauvetage (A'·B')", 20),
        ("iptm_negative", "iPTM Négatif (A_WT·B')", 20),
        ("iptm_rupture", "iPTM Rupture (A'·B_WT)", 20),
        ("f_ortho", "Score F_ortho", 16),
        ("dockq", "Score DockQ", 14),
        ("dockq_quality", "Qualité CAPRI", 16),
        ("fnat", "F_nat Contacts", 14),
        ("plddt_rescue", "pLDDT Complexe (%)", 18),
        ("is_orthogonal", "Verdict Orthogonal", 22),
    ]

    header_fill = PatternFill(start_color="2B4C7E", end_color="2B4C7E", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    ws1.row_dimensions[4].height = 28
    for col_idx, (key, label, width) in enumerate(cols_summary, start=1):
        cell = ws1.cell(row=4, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width

    # Row Styling Fills
    fill_champ = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")  # Vert clair
    fill_pass = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_regular = Font(name="Segoe UI", size=10)
    font_code = Font(name="Consolas", size=9, bold=True)

    for row_idx, row in df.iterrows():
        r_num = row_idx + 5
        ws1.row_dimensions[r_num].height = 20
        
        f_ortho_val = row.get('f_ortho', 0.0)
        is_champ = f_ortho_val >= 0.20
        is_ortho = f_ortho_val > 0.0
        
        row_fill = fill_champ if is_champ else (fill_pass if is_ortho else (fill_alt if row_idx % 2 == 1 else fill_white))

        dockq_val = row.get('dockq', 0.0)
        dockq_q = row.get('dockq_quality', 'N/A')
        fnat_val = row.get('fnat', 0.0)
        plddt_val = row.get('plddt_rescue', row.get('plddt_protein', 0.0))

        status_text = "CHAMPION ORTHOGONAL" if is_champ else ("ORTHOGONAL" if is_ortho else "NON ORTHOGONAL")

        row_values = [
            row.get('design_id', ''),
            row.get('parent_a_motif', ''),
            row.get('iptm_rescue', 0.0),
            row.get('iptm_negative', 0.0),
            row.get('iptm_rupture', 0.0),
            f_ortho_val,
            dockq_val if pd.notnull(dockq_val) else 0.0,
            dockq_q if pd.notnull(dockq_q) else 'N/A',
            fnat_val if pd.notnull(fnat_val) else 0.0,
            plddt_val if pd.notnull(plddt_val) else 0.0,
            status_text
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws1.cell(row=r_num, column=col_idx, value=val)
            cell.fill = row_fill
            cell.border = thin_border
            cell.font = font_code if col_idx == 1 else (font_bold if col_idx in [3, 6, 11] else font_regular)
            
            if col_idx in [3, 4, 5, 6, 7, 9, 10]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if col_idx in [3, 4, 5, 6, 7]:
                    cell.number_format = "0.0000"
                elif col_idx == 9:
                    cell.number_format = "0.000"
                elif col_idx == 10:
                    cell.number_format = "0.0"
            else:
                cell.alignment = Alignment(horizontal="center" if col_idx in [2, 8, 11] else "left", vertical="center")

    # -------------------------------------------------------------
    # Sheet 2: Raw High-Resolution Metrics
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Raw All-Atom Metrics")
    ws2.views.sheetView[0].showGridLines = True

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = max(len(col_name) + 4, 14)

    for row_idx, row in df.iterrows():
        for col_idx, col_name in enumerate(df.columns, start=1):
            val = row[col_name]
            cell = ws2.cell(row=row_idx + 2, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            if isinstance(val, float):
                cell.number_format = "0.0000"

    wb.save(xlsx_path)
    print(f"[Styled Excel Export] Saved clean interactive spreadsheet -> {xlsx_path}")

def generate_interactive_html(csv_path, html_path):
    if not os.path.exists(csv_path):
        return

    df = pd.read_csv(csv_path)

    rows_html = []
    for idx, row in df.iterrows():
        iptm_r = float(row.get('iptm_rescue', 0.0))
        iptm_neg = float(row.get('iptm_negative', 0.0))
        iptm_rup = float(row.get('iptm_rupture', 0.0))
        f_ortho = float(row.get('f_ortho', 0.0))
        dockq_val = row.get('dockq', 0.0)
        dockq_q = row.get('dockq_quality', 'N/A')
        f_nat = float(row.get('fnat', 0.0)) if pd.notnull(row.get('fnat')) else 0.0
        plddt = float(row.get('plddt_rescue', row.get('plddt_protein', 0.0)))

        is_champ = f_ortho >= 0.20
        is_ortho = f_ortho > 0.0
        
        if is_champ:
            status_badge = '<span class="badge badge-champion">Champion Orthogonal</span>'
        elif is_ortho:
            status_badge = '<span class="badge badge-success">Orthogonal</span>'
        else:
            status_badge = '<span class="badge badge-negative">Non Orthogonal</span>'

        row_str = f"""
        <tr>
            <td style="font-weight: 700; color: var(--text-muted);">{idx+1}</td>
            <td><code>{row.get('design_id', '')}</code></td>
            <td><span class="strategy-tag">{row.get('parent_a_motif', 'Co-adaptation')}</span></td>
            <td>
                <div class="metric-val" style="color: #10b981;">{iptm_r:.4f}</div>
                <div class="progress-bar-bg"><div class="progress-bar-fill green" style="width: {min(iptm_r*100, 100):.1f}%"></div></div>
            </td>
            <td>
                <div class="metric-val" style="color: #38bdf8;">{iptm_neg:.4f}</div>
                <div class="progress-bar-bg"><div class="progress-bar-fill blue" style="width: {min(iptm_neg*100, 100):.1f}%"></div></div>
            </td>
            <td>
                <div class="metric-val" style="color: #64748b;">{iptm_rup:.4f}</div>
            </td>
            <td style="font-weight: 700; font-size: 14px; color: {'#10b981' if f_ortho >= 0.20 else ('#38bdf8' if f_ortho > 0 else '#ef4444')};">{f_ortho:+.4f}</td>
            <td><b>{dockq_val if pd.notnull(dockq_val) else 0.0:.4f}</b> <small style="color: var(--text-muted);">({dockq_q})</small></td>
            <td>{f_nat:.3f}</td>
            <td><b>{plddt:.1f}%</b></td>
            <td>{status_badge}</td>
        </tr>
        """
        rows_html.append(row_str)

    max_iptm_r = df['iptm_rescue'].max() if not df.empty else 0.0
    max_f_ortho = df['f_ortho'].max() if not df.empty else 0.0
    total_designs = len(df)
    ortho_count = len(df[df['f_ortho'] > 0])

    html_content = f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>OrthoPPInterface - Tableau de Bord Interactif</title>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
    <style>
        :root {{
            --bg: #0f172a;
            --card-bg: #1e293b;
            --border: #334155;
            --text: #f8fafc;
            --text-muted: #94a3b8;
            --accent: #38bdf8;
            --green: #10b981;
            --blue: #3b82f6;
            --yellow: #f59e0b;
        }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{
            font-family: 'Inter', sans-serif;
            background-color: var(--bg);
            color: var(--text);
            padding: 30px;
            line-height: 1.5;
        }}
        .header {{
            background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
            border: 1px solid var(--border);
            border-radius: 16px;
            padding: 28px;
            margin-bottom: 24px;
            box-shadow: 0 4px 20px rgba(0, 0, 0, 0.3);
        }}
        .header h1 {{
            font-size: 24px;
            font-weight: 700;
            color: var(--text);
            display: flex;
            align-items: center;
            gap: 12px;
        }}
        .header p {{
            color: var(--text-muted);
            margin-top: 6px;
            font-size: 14px;
        }}
        .stats-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
            gap: 16px;
            margin-bottom: 24px;
        }}
        .stat-card {{
            background-color: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 12px;
            padding: 20px;
            position: relative;
        }}
        .stat-card .label {{
            font-size: 12px;
            text-transform: uppercase;
            letter-spacing: 0.05em;
            color: var(--text-muted);
            font-weight: 600;
        }}
        .stat-card .val {{
            font-size: 26px;
            font-weight: 700;
            margin-top: 6px;
            font-family: 'JetBrains Mono', monospace;
        }}
        .search-container {{
            margin-bottom: 16px;
            display: flex;
            gap: 12px;
        }}
        .search-input {{
            flex: 1;
            background-color: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 8px;
            padding: 10px 16px;
            color: var(--text);
            font-size: 14px;
            outline: none;
        }}
        .search-input:focus {{ border-color: var(--accent); }}
        .table-container {{
            background-color: var(--card-bg);
            border: 1px solid var(--border);
            border-radius: 12px;
            overflow-x: auto;
            box-shadow: 0 4px 20px rgba(0, 0, 0, 0.2);
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            text-align: left;
            font-size: 13px;
        }}
        th {{
            background-color: #0f172a;
            color: var(--text-muted);
            font-weight: 600;
            padding: 14px 16px;
            border-bottom: 1px solid var(--border);
            text-transform: uppercase;
            font-size: 11px;
            letter-spacing: 0.05em;
            cursor: pointer;
            user-select: none;
        }}
        th:hover {{ color: var(--accent); }}
        td {{
            padding: 14px 16px;
            border-bottom: 1px solid var(--border);
            color: var(--text);
        }}
        tr:hover td {{ background-color: rgba(255, 255, 255, 0.02); }}
        code {{
            font-family: 'JetBrains Mono', monospace;
            background-color: #0f172a;
            padding: 4px 8px;
            border-radius: 6px;
            color: #38bdf8;
            font-size: 12px;
        }}
        .strategy-tag {{
            background-color: rgba(56, 189, 248, 0.1);
            color: #38bdf8;
            border: 1px solid rgba(56, 189, 248, 0.2);
            padding: 4px 10px;
            border-radius: 20px;
            font-size: 11px;
            font-weight: 500;
            white-space: nowrap;
        }}
        .metric-val {{ font-weight: 600; font-family: 'JetBrains Mono', monospace; font-size: 13px; }}
        .progress-bar-bg {{
            background-color: #334155;
            height: 6px;
            border-radius: 3px;
            width: 100%;
            margin-top: 6px;
            overflow: hidden;
        }}
        .progress-bar-fill.green {{ background-color: #10b981; height: 100%; }}
        .progress-bar-fill.blue {{ background-color: #3b82f6; height: 100%; }}
        .badge {{
            padding: 4px 10px;
            border-radius: 20px;
            font-size: 11px;
            font-weight: 600;
            display: inline-block;
            white-space: nowrap;
        }}
        .badge-champion {{ background-color: rgba(16, 185, 129, 0.2); color: #10b981; border: 1px solid rgba(16, 185, 129, 0.4); }}
        .badge-success {{ background-color: rgba(56, 189, 248, 0.15); color: #38bdf8; border: 1px solid rgba(56, 189, 248, 0.3); }}
        .badge-negative {{ background-color: rgba(239, 68, 68, 0.15); color: #f87171; border: 1px solid rgba(239, 68, 68, 0.3); }}
    </style>
</head>
<body>
    <div class="header">
        <h1>OrthoPPInterface - Matrice d'Orthogonalité Protéine-Protéine</h1>
        <p>Conception Rationnelle & Évaluation RoseTTAFold-3 All-Atom (RF3) GPU | Conformité CAPRI DockQ</p>
    </div>

    <div class="stats-grid">
        <div class="stat-card">
            <div class="label">Paires Évaluées</div>
            <div class="val" style="color: var(--accent);">{total_designs}</div>
        </div>
        <div class="stat-card">
            <div class="label">Paires Orthogonales Valides</div>
            <div class="val" style="color: var(--green);">{ortho_count} <small style="font-size: 14px; color: var(--text-muted);">({ortho_count/total_designs*100:.1f}%)</small></div>
        </div>
        <div class="stat-card">
            <div class="label">Affinité Sauvetage Max (iPTM)</div>
            <div class="val" style="color: var(--green);">{max_iptm_r:.4f}</div>
        </div>
        <div class="stat-card">
            <div class="label">Score F_ortho Maximal</div>
            <div class="val" style="color: var(--green);">+{max_f_ortho:.4f}</div>
        </div>
    </div>

    <div class="search-container">
        <input type="text" id="searchInput" class="search-input" placeholder="Filtrer par nom de candidat ou motif A'..." onkeyup="filterTable()">
    </div>

    <div class="table-container">
        <table id="orthoTable">
            <thead>
                <tr>
                    <th onclick="sortTable(0)">#</th>
                    <th onclick="sortTable(1)">Paire Orthogonale (A'·B')</th>
                    <th onclick="sortTable(2)">Motif A' Parent</th>
                    <th onclick="sortTable(3)">iPTM Sauvetage (A'·B')</th>
                    <th onclick="sortTable(4)">iPTM Négatif (A_WT·B')</th>
                    <th onclick="sortTable(5)">iPTM Rupture (A'·B_WT)</th>
                    <th onclick="sortTable(6)">Score F_ortho</th>
                    <th onclick="sortTable(7)">Score DockQ</th>
                    <th onclick="sortTable(8)">F_nat</th>
                    <th onclick="sortTable(9)">pLDDT Complexe</th>
                    <th onclick="sortTable(10)">Verdict</th>
                </tr>
            </thead>
            <tbody>
                {''.join(rows_html)}
            </tbody>
        </table>
    </div>

    <script>
        function filterTable() {{
            const input = document.getElementById('searchInput');
            const filter = input.value.toUpperCase();
            const table = document.getElementById('orthoTable');
            const tr = table.getElementsByTagName('tr');
            for (let i = 1; i < tr.length; i++) {{
                let text = tr[i].textContent || tr[i].innerText;
                tr[i].style.display = text.toUpperCase().indexOf(filter) > -1 ? "" : "none";
            }}
        }}

        function sortTable(n) {{
            const table = document.getElementById("orthoTable");
            let rows, switching, i, x, y, shouldSwitch, dir, switchcount = 0;
            switching = true;
            dir = "desc";
            while (switching) {{
                switching = false;
                rows = table.rows;
                for (i = 1; i < (rows.length - 1); i++) {{
                    shouldSwitch = false;
                    x = rows[i].getElementsByTagName("TD")[n];
                    y = rows[i + 1].getElementsByTagName("TD")[n];
                    let xVal = parseFloat(x.innerText) || x.innerText.toLowerCase();
                    let yVal = parseFloat(y.innerText) || y.innerText.toLowerCase();
                    if (dir === "asc") {{
                        if (xVal > yVal) {{ shouldSwitch = true; break; }}
                    }} else if (dir === "desc") {{
                        if (xVal < yVal) {{ shouldSwitch = true; break; }}
                    }}
                }}
                if (shouldSwitch) {{
                    rows[i].parentNode.insertBefore(rows[i + 1], rows[i]);
                    switching = true;
                    switchcount++;
                }} else {{
                    if (switchcount === 0 && dir === "desc") {{
                        dir = "asc";
                        switching = true;
                    }}
                }}
            }}
        }}
    </script>
</body>
</html>
"""
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"[Interactive HTML Export] Saved clean visual dashboard -> {html_path}")
